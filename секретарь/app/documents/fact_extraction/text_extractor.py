"""
ЭТАП 3.4–3.6
Извлечение текста из документов.
Приоритет:
PDF text layer;
OCR, если text layer отсутствует/непригоден;
изображения → OCR;
DOCX/XLSX → структурированное чтение.
Оригинальные файлы НЕ изменяются.
"""
import os
from dataclasses import dataclass, field
from typing import List, Tuple
import numpy as np
import pypdfium2 as pdfium
try:
    from rapidocr_onnxruntime import RapidOCR
except Exception:
    RapidOCR = None
try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None
try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None


@dataclass
class TextExtractionResult:
    lines: List[Tuple[str, float, int]] = field(default_factory=list)

    extraction_method: str = "UNKNOWN"

    pages_total: int = 0
    pages_read: int = 0
    pages_failed: int = 0

    error: str = ""
    ocr_version: str = "rapidocr-onnxruntime"

    technical_quality: str = "GOOD"

    handwriting_detected: bool = False
    unreadable: bool = False

    # ========================================================
    # ПУНКТ 3.5 — РЕЗУЛЬТАТ ЧТЕНИЯ ХРАНИТСЯ ОТДЕЛЬНО
    # ========================================================

    extracted_text: str = ""

    ocr_result: List[dict] = field(
        default_factory=list
    )


class TextExtractor:
    def __init__(self):
        self.ocr = None
        if RapidOCR is not None:
            self.ocr = RapidOCR()

    # ============================================================
    # PUBLIC
    # ============================================================
    def extract(self, filepath: str) -> TextExtractionResult:
        extension = os.path.splitext(filepath)[1].lower()
        if extension == ".pdf":
            return self._extract_from_pdf(filepath)
        if extension in {".jpg", ".jpeg", ".png"}:
            return self._extract_from_image(filepath)
        if extension == ".docx":
            return self._extract_from_docx(filepath)
        if extension == ".xlsx":
            return self._extract_from_xlsx(filepath)
        if extension in {".doc", ".xls"}:
            return TextExtractionResult(
                extraction_method="UNSUPPORTED_FORMAT",
                technical_quality="REQUIRES_REVIEW",
                unreadable=True,
                error=(
                    "Формат поддержан архитектурой, "
                    "но для чтения нужен конвертер: "
                    f"{extension}"
                ),
            )
        return TextExtractionResult(
            extraction_method="UNSUPPORTED_FORMAT",
            technical_quality="REQUIRES_REVIEW",
            unreadable=True,
            error=f"Неподдерживаемое расширение: {extension}",
        )

    # ============================================================
    # PDF
    # ============================================================
    def _extract_from_pdf(self, filepath: str):
        result = TextExtractionResult(
            extraction_method="TEXT_LAYER"
        )
        try:
            pdf = pdfium.PdfDocument(filepath)
            result.pages_total = len(pdf)
            text_lines = []
            for page_index in range(len(pdf)):
                try:
                    page = pdf[page_index]
                    textpage = page.get_textpage()
                    text = textpage.get_text_range()
                    text = (text or "").strip()
                    if len(text) >= 50:
                        for line in text.splitlines():
                            line = line.strip()
                            if line:
                                text_lines.append(
                                    (
                                        line,
                                        1.0,
                                        page_index + 1,
                                    )
                                )
                        result.pages_read += 1
                    else:
                        raise ValueError(
                            "Недостаточный text layer"
                        )
                except Exception:
                    # Для конкретной страницы переходим к OCR.
                    try:
                        page = pdf[page_index]
                        bitmap = page.render(
                            scale=2.0
                        )
                        image = bitmap.to_numpy()
                        ocr_lines = self._run_ocr(
                            image,
                            page_index + 1,
                        )
                        if ocr_lines:
                            text_lines.extend(ocr_lines)
                        result.pages_read += 1
                    except Exception:
                        result.pages_failed += 1
            result.lines = text_lines

            # ========================================================
            # ПУНКТ 3.5 — ОТДЕЛЬНЫЙ РЕЗУЛЬТАТ ЧТЕНИЯ
            # ========================================================

            result.extracted_text = "\n".join(
                text
                for text, _, _ in text_lines
            )

            result.ocr_result = [
                {
                    "text": text,
                    "confidence": confidence,
                    "page": page,
                }
                for text, confidence, page in text_lines
                if confidence < 0.99
            ]

            text_pages = set(
                page
                for _, _, page in text_lines
            )
            if result.pages_failed > 0:
                result.technical_quality = "REQUIRES_REVIEW"
            if not text_lines:
                result.unreadable = True
                result.technical_quality = "REQUIRES_REVIEW"
                result.extraction_method = "OCR"
                result.error = (
                    "Не удалось извлечь текст "
                    "ни text layer, ни OCR."
                )
            elif result.pages_failed == 0:
                # Определяем метод по confidence.
                if any(
                    confidence < 0.99
                    for _, confidence, _ in text_lines
                ):
                    result.extraction_method = "MIXED"
                else:
                    result.extraction_method = "TEXT_LAYER"
            return result
        except Exception as exc:
            result.extraction_method = "READ_ERROR"
            result.technical_quality = "REQUIRES_REVIEW"
            result.unreadable = True
            result.error = f"{type(exc).__name__}: {exc}"
            return result

    # ============================================================
    # IMAGE
    # ============================================================
    def _extract_from_image(self, filepath: str):
        result = TextExtractionResult(
            extraction_method="OCR",
            pages_total=1,
        )
        try:
            from PIL import Image
            image = np.array(Image.open(filepath))
            lines = self._run_ocr(image, 1)
            result.lines = lines

            # ========================================================
            # ПУНКТ 3.5 — ОТДЕЛЬНЫЙ OCR-РЕЗУЛЬТАТ
            # ========================================================

            result.extracted_text = "\n".join(
                text
                for text, _, _ in lines
            )

            result.ocr_result = [
                {
                    "text": text,
                    "confidence": confidence,
                    "page": page,
                }
                for text, confidence, page in lines
            ]

            result.pages_read = 1 if lines else 0
            if not lines:
                result.unreadable = True
                result.technical_quality = "REQUIRES_REVIEW"
                result.error = "OCR не вернул текст."
            return result
        except Exception as exc:
            result.pages_failed = 1
            result.unreadable = True
            result.technical_quality = "REQUIRES_REVIEW"
            result.error = f"{type(exc).__name__}: {exc}"
            return result

    # ============================================================
    # OCR
    # ============================================================
    def _run_ocr(self, image, page_number: int):
        if self.ocr is None:
            return []
        try:
            output, _ = self.ocr(image)
        except Exception:
            return []
        lines = []
        if output is None:
            return lines
        for item in output:
            try:
                # RapidOCR обычно:
                # [box, text, confidence]
                text = str(item[1]).strip()
                confidence = float(item[2])
                if text:
                    lines.append(
                        (
                            text,
                            confidence,
                            page_number,
                        )
                    )
            except Exception:
                continue
        return lines

    # ============================================================
    # DOCX
    # ============================================================
    def _extract_from_docx(self, filepath: str):
        result = TextExtractionResult(
            extraction_method="DOCX_TEXT",
            pages_total=1,
        )
        if DocxDocument is None:
            result.unreadable = True
            result.technical_quality = "REQUIRES_REVIEW"
            result.error = "python-docx не установлен."
            return result
        try:
            document = DocxDocument(filepath)
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if text:
                    result.lines.append(
                        (
                            text,
                            1.0,
                            1,
                        )
                    )
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if text:
                            result.lines.append(
                                (
                                    text,
                                    1.0,
                                    1,
                                )
                            )

            # ========================================================
            # ПУНКТ 3.5 — ОТДЕЛЬНЫЙ ТЕКСТ ДОКУМЕНТА
            # ========================================================

            result.extracted_text = "\n".join(
                text
                for text, _, _ in result.lines
            )

            result.pages_read = 1
            if not result.lines:
                result.unreadable = True
                result.technical_quality = "REQUIRES_REVIEW"
                result.error = "DOCX не содержит текста."
            return result
        except Exception as exc:
            result.unreadable = True
            result.technical_quality = "REQUIRES_REVIEW"
            result.error = f"{type(exc).__name__}: {exc}"
            return result

    # ============================================================
    # XLSX
    # ============================================================
    def _extract_from_xlsx(self, filepath: str):
        result = TextExtractionResult(
            extraction_method="XLSX_TEXT"
        )
        if load_workbook is None:
            result.unreadable = True
            result.technical_quality = "REQUIRES_REVIEW"
            result.error = "openpyxl не установлен."
            return result
        try:
            workbook = load_workbook(
                filepath,
                read_only=True,
                data_only=True,
            )
            for sheet_index, sheet in enumerate(
                workbook.worksheets,
                start=1,
            ):
                result.pages_total += 1
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        text = str(cell.value).strip()
                        if text:
                            result.lines.append(
                                (
                                    text,
                                    1.0,
                                    sheet_index,
                                )
                            )
                result.pages_read += 1

            # ========================================================
            # ПУНКТ 3.5 — ОТДЕЛЬНЫЙ ТЕКСТ ТАБЛИЦЫ
            # ========================================================

            result.extracted_text = "\n".join(
                text
                for text, _, _ in result.lines
            )

            if not result.lines:
                result.unreadable = True
                result.technical_quality = "REQUIRES_REVIEW"
                result.error = "XLSX не содержит данных."
            return result
        except Exception as exc:
            result.unreadable = True
            result.technical_quality = "REQUIRES_REVIEW"
            result.error = f"{type(exc).__name__}: {exc}"
            return result
